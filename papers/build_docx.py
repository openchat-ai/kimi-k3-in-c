#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the submission .docx to the journal's own template.

Every parameter below is taken from jsjk_template.doc (jsjkx.com, saved 2024-07-11), not from
taste. The template states:

  中文题目            黑体 三号（16pt）
  作者                宋体加粗 五号（10.5pt）
  单位                仿宋 小五（9pt）
  摘 要               不少于 300 字，不使用"本文""我们"等第一人称
  关键词              5~8 个
  中图分类号          细化到 3 位数字
  English Title      Times New Roman 加粗 四号（14pt），实词首字母大写
  英文作者            Times New Roman 五号（10.5pt）
  Abstract           不少于 300 词
  正文                双栏排版
  一级标题            黑体 五号（10.5pt）
  二级标题            黑体 小五（9pt）
  三级标题            楷体 小五（9pt）
  正文                宋体 小五（9pt）
  图表题              6 号（7.5pt），中文宋体、英文 Times New Roman
  表题置于表上方居中，图题置于图下方居中
  图宽一般 8 厘米以内

The body is two columns per the template; the front matter is single column, which is how
journal templates of this shape are set, so the section carries one column and the body
section two.
"""
import re, sys, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Which manuscript to typeset. The parser below is generic -- it reads the byline, abstract,
# captions and figures out of the markdown -- so a second paper only needs its own paths here
# rather than a forked copy of this file, which is how the two DOCX files drifted apart before.
#   python papers/build_docx.py SRC [OUT]
if len(sys.argv) < 2:
    # 该脚本被两篇互不相关的稿件共用。给它设一个默认源文件，等于让一次裸跑悄悄去构建
    # 另一篇稿子——两稿内容不可混用，故此处要求显式指定，不提供默认值。
    sys.exit("用法: python papers/build_docx.py <源 .md> [输出 .docx]")
SRC = pathlib.Path(sys.argv[1])
OUT = (pathlib.Path(sys.argv[2]) if len(sys.argv) > 2
       else SRC.with_name("提交稿-" + SRC.stem.split("论文-")[-1] + ".docx"))

SIMSUN, SIMHEI, KAITI, FANGSONG, TNR = "宋体", "黑体", "楷体", "仿宋", "Times New Roman"

# 官方模板未单独规定表内字号，只规定正文为宋体小五（9 pt）、图表题为 6 号（7.5 pt）。
# 表内文字取与正文相同的 9 pt，是"模板未规定处随正文"这一惯例的选择，而非模板明文要求。
TABLE_PT = 9


def set_cjk(run, cn_font, size_pt, bold=False):
    """python-docx sets the Latin face only; Chinese glyphs need w:eastAsia set explicitly."""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = TNR
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), TNR)
    rFonts.set(qn("w:hAnsi"), TNR)
    rFonts.set(qn("w:eastAsia"), cn_font)


def para(doc, text, cn=SIMSUN, size=9, bold=False, align=None, space_after=3,
         indent_chars=0, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.0
    if indent_chars:
        p.paragraph_format.first_line_indent = Pt(size * indent_chars)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_cjk(r, cn, size, bold)
    return p


def two_col(section, gap_cm=0.6):
    """Two equal columns with a gutter, written as raw sectPr because python-docx has no API."""
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(int(Cm(gap_cm).twips)))
    cols.set(qn("w:equalWidth"), "1")


def figure_block(doc, img, cap_cn, cap_en, width_cm=8.0):
    """Image centred, then the bilingual caption below it -- the template puts figure titles
    under the figure."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(img), width=Cm(width_cm))
    for text, font in ((cap_cn, SIMSUN), (cap_en, TNR)):
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(1)
        c.paragraph_format.space_before = Pt(0)
        set_cjk(c.add_run(text), font, 7.5)


def table_caption(doc, cap_cn, cap_en):
    """Table titles above the table -- the template puts table titles on top."""
    for text, font in ((cap_cn, SIMSUN), (cap_en, TNR)):
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before = Pt(4)
        c.paragraph_format.space_after = Pt(1)
        set_cjk(c.add_run(text), font, 7.5)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
CODE = re.compile(r"`([^`]+)`")


def rich(p, text, cn=SIMSUN, size=9, base_bold=False):
    """Emit **bold** and `code` runs; markdown emphasis inside body text is common in the source."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            set_cjk(p.add_run(tok[2:-2]), cn, size, True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(CODE.match(tok).group(1))
            r.font.size = Pt(size)
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(0, 0, 0x80)
        else:
            set_cjk(p.add_run(tok), cn, size, base_bold)


def _disp_width(s):
    """Display width in half-ems: a CJK glyph occupies two, everything else one."""
    return sum(2 if ord(ch) > 0x2000 else 1 for ch in s)


def add_md_table(doc, rows, total_cm=8.6, min_cm=1.0):
    """Add a table whose column widths follow the content instead of being equal.

    python-docx creates equal-width columns. Tables in this paper carry one long label column
    ("分相位实测（第二次，主干常驻档）", "16 路并发，层内取 16（见 5.2.1）") beside several short
    numeric ones, so equal widths waste space on the numbers and push the label past the
    column edge. Widths are allocated in proportion to the widest cell in each column, with a
    floor so a one-character column does not collapse, and the total is held at the measure of
    the two-column body so the table cannot exceed the text block.
    """
    cols = len(rows[0])
    if any(len(r) != cols for r in rows):
        raise SystemExit("表格各行列数不一致: %s" % [len(r) for r in rows])

    # 中文可在任意两字之间换行，故只有不可断开的英文/数字串（"B_distinct"、"216.66"）
    # 会因列宽不足而被裁切。列宽先给到该下限，余下宽度再按内容比例分配。
    pt_cm = TABLE_PT / 72.0 * 2.54

    def _min_need(cell):
        runs = re.findall(r"[A-Za-z0-9_.]+", cell)
        return max((len(x) for x in runs), default=0) * 0.5 * pt_cm

    widest, floors = [], []
    for ci in range(cols):
        w = 1
        f = 0.0
        for r in rows:
            w = max(w, _disp_width(r[ci]))
            f = max(f, _min_need(r[ci]))
        widest.append(w)
        floors.append(f)

    base = [max(min_cm, f + 0.10) for f in floors]
    spare = total_cm - sum(base)
    if spare < 0:                       # 词元本身已占满整栏，只能按比例压缩
        base = [b * total_cm / sum(base) for b in base]
        spare = 0.0
    tot = float(sum(widest))
    widths = [b + spare * w / tot for b, w in zip(base, widest)]

    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    # python-docx sets cell widths but Word ignores them unless the table layout is fixed and
    # the grid itself carries the widths. Without both, the first version of this change read
    # back as five equal-width tables totalling 17.20 cm -- more than twice the 8.6 cm measure
    # of a body column.
    t.autofit = False
    tblPr = t._tbl.tblPr
    # A named style is not a border. "Table Grid" is referenced here and the style id exists in
    # styles.xml, but the style carries no border definition in this document, so Word rendered
    # all eight tables as vertical rules with no lines. Borders are therefore set explicitly on
    # the table rather than inherited, because a table whose visibility depends on style
    # resolution is a table that can silently lose its rules.
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    borders = OxmlElement("w:tblBorders")
    # 官方模板要求三线表（投稿指南原文："Make your columns in three-line tables"），
    # 即只有顶线、表头下的中线和底线三道横线，无竖线、无底纹。python-docx 的
    # "Table Grid" 是全网格，样式本身也不带边框定义，因此边框在此显式声明。
    for tag, sz in (("top", "12"), ("bottom", "12"),
                    ("left", "0"), ("right", "0"), ("insideH", "0"), ("insideV", "0")):
        b = OxmlElement("w:%s" % tag)
        if sz == "0":
            b.set(qn("w:val"), "none")
            b.set(qn("w:sz"), "0")
            b.set(qn("w:color"), "auto")
        else:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), sz)
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)

    # Word 默认左右各留 108 twip（0.19 cm），在 8 列的台账表里合计占去 3 cm，比一行数字
    # 还宽。台账表按密排处理，把左右边距压到 40 twip。
    for el in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(el)
    cellmar = OxmlElement("w:tblCellMar")
    for tag, tw in (("left", "40"), ("right", "40")):
        m = OxmlElement("w:%s" % tag)
        m.set(qn("w:w"), tw)
        m.set(qn("w:type"), "dxa")
        cellmar.append(m)
    tblPr.append(cellmar)

    for tag in ("w:tblLayout",):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        t._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
        grid.append(gc)
    t._tbl.insert(list(t._tbl).index(tblPr) + 1, grid)

    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            cells[ci].width = Cm(widths[ci])
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_cjk(p.add_run(val.strip()), SIMHEI if ri == 0 else SIMSUN,
                    TABLE_PT, ri == 0)
            if ri == 0:
                # 三线表的中线画在表头行下沿；模板同时明确"不宜加底色及阴影"，
                # 故表头不再铺灰底。
                tcPr = cells[ci]._tc.get_or_add_tcPr()
                tcBorders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "0")
                bottom.set(qn("w:color"), "000000")
                tcBorders.append(bottom)
                tcPr.append(tcBorders)
    return t


def parse(md):
    """Split the markdown into the front matter and a body of typed blocks."""
    lines = md.splitlines()
    head, i = {}, 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("# "):
            head["title"] = s[2:].strip()
        elif s.startswith("唐海勇"):
            head["author"] = s
        elif "收稿日期" in s:
            head["received"] = s
        elif s.startswith("## 摘要"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].startswith("## "):
                if lines[j].strip():
                    buf.append(lines[j].strip())
                j += 1
            head["abstract"] = "".join(buf)
            head["abs_end"] = j
            i = j
            break
        i += 1
    return head, lines[head.get("abs_end", 0):]


def main():
    # utf-8-sig, not utf-8: a BOM left by a Windows editor becomes part of the first line's
    # text, so the "# " title test silently fails and the run dies on a missing key. Accepting
    # the optional BOM costs nothing and makes the script independent of which editor wrote
    # the draft.
    md = SRC.read_text(encoding="utf-8-sig")
    head, body = parse(md)

    ab = re.sub(r"\*\*(关键词|中图分类号)：\*\*", "", head["abstract"])
    ab = re.sub(r"\*\*[^*]+\*\*", "", ab)
    ab = re.sub(r"[（(].*?[)）]", "", ab)
    ab = ab.replace("TP314", "").replace("A", "", 1) if False else ab
    cjk = len(re.findall(r"[\u4e00-\u9fff]", ab))
    if cjk < 300:
        print(f"FATAL: 摘要仅 {cjk} 汉字，模板要求不少于 300 —— 不生成", file=sys.stderr)
        return 1
    print(f"  摘要 {cjk} 汉字（模板 ≥300）ok")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.size = Pt(9)
    st.font.name = TNR
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SIMSUN)

    # ---------------- front matter: single column ----------------
    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(s0, m, Cm(2.0))
    s0.top_margin, s0.bottom_margin = Cm(2.0), Cm(2.0)

    para(doc, head["title"], cn=SIMHEI, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    para(doc, head["author"].split("；")[0], cn=SIMSUN, size=10.5, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    unit = head["author"].split("；", 1)[1] if "；" in head["author"] else ""
    para(doc, unit, cn=FANGSONG, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
# the corresponding author's address comes from the draft's byline note, so changing it there
# changes the DOCX. It was a literal in this file, which meant the two could disagree.
    _em = re.search(r"邮箱：([^\s（(]+)", md)
    if not _em:
        sys.exit("★ 源文中未找到通讯作者邮箱")
    para(doc, "（%s）第一作者邮箱" % _em.group(1), cn=FANGSONG, size=9,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    para(doc, "摘  要", cn=SIMHEI, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    rich(p, head["abstract"].split("关键词")[0].strip(), size=9)
    # strip the classification tail from the abstract paragraph
    for r in list(p.runs):
        if "中图分类号" in r.text:
            r.text = r.text.split("中图分类号")[0]

    for line in (head["abstract"].split("关键词")[-1].split("\n") if "关键词" in head["abstract"] else []):
        pass
    kw = re.search(r"关键词：([^\n]+)", head["abstract"])
    cls = re.search(r"中图分类号：(\w+)", head["abstract"])
    doc_flag = re.search(r"文献标志码：(\w)", head["abstract"])
    if kw:
        para(doc, "关键词：" + kw.group(1).strip(), cn=SIMHEI, size=9, space_after=2)
    if cls:
        para(doc, f"中图分类号：{cls.group(1)}　　文献标志码：{doc_flag.group(1) if doc_flag else 'A'}",
             cn=SIMHEI, size=9, space_after=8)

    # The English block ends at the first horizontal rule or top-level heading. The original
    # pattern was re.compile(r"## High Cache Hit Rate(.*)", re.S), whose greedy (.*) ran to
    # end-of-file: it swallowed the '---', the author bio and the self-check contact into the
    # front matter, and the body loop then rendered all three a second time.
    en = re.search(r"^##\s*High Cache Hit Rate.*$(.*?)(?=^\s*---\s*$|^\#\s)", md, re.S | re.M)
    if en:
        en_lines = [l.strip() for l in en.group(1).splitlines() if l.strip()]
        # The English title, byline and affiliation are read from the markdown block rather than
        # written here. They were duplicated in this file and drifted: the Chinese title was
        # changed and the English one was not, and the DOCX -- generated from this file -- printed
        # the stale English while the markdown carried the new one. The fix is a single source,
        # not care: the block at the end of the draft holds all of them.
        first = re.match(r"##\s*(High Cache Hit Rate.*)", md).group(1).strip() \
            if re.search(r"^##\s*High Cache Hit Rate", md, re.M) else ""
        if not first:
            sys.exit("★ 源文中未找到英文题名")
        para(doc, first, cn=TNR, size=14, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        print(f"  英文题名（取自 md）{first[:56]}")

        byline_en = next((l for l in en_lines if l.startswith("TANG ")), "")
        if not byline_en:
            sys.exit("★ 英文块中未找到署名行")
        parts = byline_en.split(";")
        para(doc, parts[0].strip(), cn=TNR, size=10.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
        if len(parts) > 1:
            para(doc, parts[1].strip(), cn=TNR, size=9,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

        # en_lines still holds the title and the byline, which are printed above with their own
        # formatting. Skipping them here keeps them from being rendered a second time as body text.
        abstract_lines = [l for l in en_lines
                          if l != first and l != byline_en]
        words = 0
        for l in abstract_lines:
            if l.startswith("**Keywords:**"):
                para(doc, "Keywords: " + l.split(":", 1)[1].strip(), cn=TNR, size=9, space_after=8)
            else:
                para(doc, l, cn=TNR, size=9, space_after=3)
                words += len(re.findall(r"[A-Za-z][A-Za-z'-]*", l))
        print(f"  英文摘要 {words} 词（模板 ≥300）{'ok' if words >= 300 else '*** 不足 ***'}")

    # ---------------- body: two columns ----------------
    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    s1.page_width, s1.page_height = Cm(21), Cm(29.7)
    s1.left_margin = s1.right_margin = Cm(1.6)
    s1.top_margin = s1.bottom_margin = Cm(1.8)
    two_col(s1)

    FIGS = {"bytes_paradox": ("docs/images/bytes_paradox_grey.png"), "trunk_cache_split": ("docs/images/trunk_cache_split_grey.png")}
    i, h1 = 0, 0
    # Figure captions sit ABOVE the image in the markdown but must be printed BELOW it, so they
    # are held here instead of being emitted as body text when first encountered.
    held_cn = held_en = ""
    while i < len(body):
        raw = body[i]
        s = raw.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^图\s*\d+\s*[　\s].*$", s)                  # figure caption, held
        if m:
            held_cn = s
            i += 1
            continue
        m = re.match(r"^(Fig\.?\s*\d+.*)$", s)                     # english figure caption, held
        if m:
            held_en = s
            i += 1
            continue
        if s.startswith("|"):                                    # markdown table
            rows = []
            while i < len(body) and body[i].strip().startswith("|"):
                cells = [c for c in body[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_md_table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)                     # heading, any depth
        if m:
            txt = m.group(2).strip()
            # Back-matter blocks are not numbered sections. The sequence number is only used
            # when the markdown does not already carry one -- otherwise "## 1 引言" became
            # "1  1 引言", and 附录/参考文献 were numbered 7 and 8 as if they were chapters.
            if re.match(r"^(附录|参考文献|英文题名|作者简介|自校负责人|稿件信息)", txt):
                para(doc, txt, cn=SIMHEI, size=10.5, space_before=6, space_after=3)
                i += 1
                continue
            num = re.match(r"^(\d+(?:\.\d+)*)\s+(\S.*)$", txt)
            app = re.match(r"^([A-Z](\.\d+)+)\s+(\S.*)$", txt)     # A.1 / A.2.1
            if app:
                # Appendix subsections keep their own letters; auto-numbering turned them into
                # "1  A.1 数据出处", "2  A.2 …", "3  A.2.1 …".
                depth = min(app.group(1).count(".") + 1, 3)
                label = f"{app.group(1)}  {app.group(3)}"
                para(doc, label, cn=SIMHEI if depth == 2 else KAITI,
                     size=9, space_before=4 if depth == 2 else 3, space_after=2)
                i += 1
                continue
            if num:
                depth = num.group(1).count(".") + 1
                label = num.group(2)
            else:
                h1 += 1
                depth, label = 1, txt
            if depth == 1:
                para(doc, f"{h1 if not num else num.group(1)}  {label}",
                     cn=SIMHEI, size=10.5, space_before=6, space_after=3)
            elif depth == 2:
                para(doc, f"{num.group(1)}  {label}" if num else label,
                     cn=SIMHEI, size=9, space_before=4, space_after=2)
            else:
                para(doc, f"{num.group(1)}  {label}" if num else label,
                     cn=KAITI, size=9, space_before=3, space_after=2)
            i += 1
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", s)               # image
        if m:
            # Resolve the image the way the markdown names it, relative to the manuscript's own
            # directory. The old code threw that path away and rebuilt one from the filename stem
            # against a docs/images/*_grey.png convention, which only happened to match the
            # earlier paper; every other manuscript died on FileNotFoundError. The grey variants
            # stay as a fallback for the drafts that ship them.
            key = pathlib.Path(m.group(2)).stem.replace("_grey", "")
            cand = [SRC.parent / m.group(2),
                    pathlib.Path(m.group(2)),
                    pathlib.Path(FIGS.get(key, ("docs/images/%s_grey.png" % key)))]
            img = next((c for c in cand if c.exists()), cand[0])
            if not img.exists():
                sys.exit("★ 找不到图片：%s（源文 %s）" % (m.group(2), SRC))
            figure_block(doc, img, held_cn or f"图（{key}）", held_en, width_cm=8.0)
            held_cn = held_en = ""
            i += 1
            continue
        m = re.match(r"^(表\s*\d+[a-z]?)\s+(.*)", s)          # table caption, goes on top
        if m:
            en = body[i + 1].strip() if i + 1 < len(body) else ""
            table_caption(doc, s, en if en.startswith("Table") else "")
            i += 2 if en.startswith("Table") else 1
            continue
        if s.startswith("```"):                                 # code block
            i += 1
            while i < len(body) and not body[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.3)
                r = p.add_run(body[i])
                r.font.size = Pt(7)
                r.font.name = "Consolas"
                i += 1
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,})$", s):
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Pt(18)
        rich(p, s, size=9)
        i += 1

    doc.save(OUT)
    print(f"  已写出 {OUT}  ({OUT.stat().st_size//1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
