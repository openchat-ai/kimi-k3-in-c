#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Check that no table exceeds the measure of a body column.

The rule the user set: column widths should follow the content and must not overflow. This
reads the generated file rather than the builder's intent, because python-docx silently
ignores cell widths unless the layout is switched off, which is exactly the kind of thing
that looks right in code and wrong in the document.

Body measure: A4 21 cm, margins 1.6 cm each side, two equal columns with a 0.6 cm gutter,
so each column is (21 - 1.6 - 1.6 - 0.6) / 2 = 8.6 cm.
"""
import sys
from docx import Document
from docx.shared import Cm

doc = Document("papers/提交稿-缓存高命中与词元低输出.docx")
MEASURE = 8.6

bad = 0
for ti, t in enumerate(doc.tables, 1):
    widths = []
    for c in t.columns:
        w = c.width
        widths.append(0.0 if w is None else w.cm)
    total = sum(widths)
    head = " | ".join((c.text or "").strip()[:10] for c in t.rows[0].cells)
    flag = ""
    if any(w == 0 for w in widths):
        flag = "  *** 有列未设定宽度 ***"
        bad += 1
    elif total > MEASURE + 0.05:
        flag = "  *** 超出正文栏宽 ***"
        bad += 1
    print("  表%d  %d 列  合计 %.2f cm / %.2f cm  %s%s"
          % (ti, len(widths), total, MEASURE, head, flag))
    print("       列宽 %s" % "  ".join("%.2f" % w for w in widths))

print()
if bad:
    print("  ★ %d 张表不合格" % bad)
    sys.exit(1)
print("  全部 %d 张表在正文栏宽内，且每列均已设定宽度" % len(doc.tables))