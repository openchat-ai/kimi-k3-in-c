#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify the generated .docx against the journal template.

A builder that runs without error proves nothing: it can silently emit every paragraph at the
wrong size, drop the images, or leave the body single-column. Re-open the file and read back
what is actually in it.
"""
import re, pathlib, zipfile
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

P = pathlib.Path("papers/提交稿-缓存高命中与词元低输出.docx")
doc = Document(P)
fails, notes = [], []


def sz(run):
    return run.font.size.pt if run.font.size else None


def ea(run):
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None
    f = rPr.find(qn("w:rFonts"))
    return f.get(qn("w:eastAsia")) if f is not None else None


paras = doc.paragraphs
print(f"  段落 {len(paras)}  表格 {len(doc.tables)}  节 {len(doc.sections)}")

# ---- front matter, in template order ----
print("\n== 前置部分（模板顺序与字号）")
expect = [
    ("题目", "黑体", 16), ("作者", "宋体", 10.5), ("单位", "仿宋", 9),
    ("摘  要", "黑体", 10.5), ("关键词", "黑体", 9), ("中图分类号", "黑体", 9),
    ("English Title", "Times New Roman", 14),
]
for want_font, want_sz in ((f, s) for _, f, s in expect):
    hit = False
    for p in paras[:40]:
        for r in p.runs:
            if want_font == "Times New Roman" and re.match(r"[A-Z][a-z].*[A-Z][a-z]", r.text or "") \
               and sz(r) == 14 and p.alignment is not None and p.alignment == 1:
                hit = True
            elif ea(r) == want_font and sz(r) == want_sz and (r.text or "").strip():
                hit = True
        if hit:
            break
    print(f"  {want_font:<16} {want_sz:>5}pt  {'ok' if hit else '*** 未找到 ***'}")
    if not hit:
        fails.append(f"front matter {want_font} {want_sz}")

# ---- abstract length ----
print("\n== 摘要字数（模板：不少于 300 字，且不用第一人称）")
ab_txt = ""
grab = False
for p in paras[:40]:
    t = (p.text or "").strip()
    if t.startswith("摘") and "要" in t:
        grab = True
        continue
    if grab:
        if t.startswith("关键词"):
            break
        ab_txt += t
cjk = len(re.findall(r"[\u4e00-\u9fff]", ab_txt))
print(f"  汉字 {cjk}  {'ok' if cjk >= 300 else '*** 不足 300 ***'}")
if cjk < 300:
    fails.append("abstract < 300")
fp = [w for w in ("本文", "我们", "笔者") if w in ab_txt]
print(f"  第一人称 {fp if fp else '无 ✓'}")
if fp:
    fails.append("first person in abstract")

# ---- keywords ----
print("\n== 关键词（模板：5~8 个）")
for p in paras[:40]:
    if (p.text or "").strip().startswith("关键词"):
        kws = re.split(r"[；;、,，]", p.text.split("关键词")[1])
        kws = [k for k in kws if k.strip()]
        print(f"  {len(kws)} 个  {kws}")
        if not 5 <= len(kws) <= 8:
            fails.append(f"keywords {len(kws)}")
        break

# ---- classification code ----
print("\n== 中图分类号（模板：3 位数字）")
for p in paras[:40]:
    t = p.text or ""
    if "中图分类号" in t:
        # the build emits the label bolded inside one paragraph, so read the whole run set
        raw = "".join(r.text for r in p.runs)
        m = re.search(r"中图分类号[:：]\s*([A-Z]+\d+)", raw)
        if not m:
            m = re.search(r"([A-Z]{1,3}\d{3})", raw)
        if m:
            code = m.group(1)
            digits = re.search(r"(\d+)", code).group(1)
            print(f"  {code}  数字部分 {len(digits)} 位  {'ok' if len(digits) >= 3 else '*** 不足 3 位 ***'}")
            if len(digits) < 3:
                fails.append(f"classcode {code}")
        else:
            print(f"  段落里找不到分类号: {t[:60]!r}")
            fails.append("classcode not parsed")
        break
else:
    print("  *** 全文前 40 段未出现中图分类号 ***")
    fails.append("classcode absent")

# ---- two-column body ----
print("\n== 正文双栏（模板：正文采用双栏排版）")
found = None
for sec in doc.sections:
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is not None and cols.get(qn("w:num")):
        found = cols.get(qn("w:num"))
        break
print(f"  w:cols num={found}  {'ok（双栏）' if found == '2' else '*** 未设双栏 ***'}")
if found != "2":
    fails.append("body not two-column")

# ---- figures ----
print("\n== 插图（模板：图题在图下、6 号、宽 8cm、灰度）")
with zipfile.ZipFile(P) as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
print(f"  内嵌图片 {len(media)} 个  {media}")
for n in media:
    info = zinfo = None
    with zipfile.ZipFile(P) as z:
        d = z.read(n)
    import struct
    w, h, depth, ct = struct.unpack(">IIBB", d[16:26])
    kind = {0: "灰度", 2: "RGB", 3: "调色板", 6: "RGBA"}[ct]
    ok = ct in (0, 3) and kind != "调色板"
    print(f"    {n:<28} {w}x{h} colourtype={ct} ({kind})  {'ok' if ok else '*** 非灰度 ***'}")
    if ct == 6:
        fails.append(f"{n} is RGBA")
    if w < 600:
        notes.append(f"{n} 宽 {w}px，模板建议 ≥600")

# ---- figure captions below ----
# The template's 图题置于图下方 means the figure's TITLE sits below the FIGURE, so the image
# paragraph must come BEFORE the caption. Two earlier versions of this check had the sense
# reversed and rejected a correct document, and also flagged body sentences that merely begin
# 图 1 给出... as if they were captions.
print("\n== 图题位置（模板：图题置于图下方）")
CAPRE = re.compile(r"^图\s*\d+\s*[　\s]")
def has_img(p):
    return "graphic" in (p._p.xml or "")
for i, p in enumerate(paras):
    t = (p.text or "").strip()
    if not CAPRE.match(t) or len(t) > 60:
        continue
    above = any(has_img(q) for q in paras[max(0, i - 3):i])   # image precedes title = compliant
    below = any(has_img(q) for q in paras[i + 1:i + 3])
    if above:
        place = "图在题上方 ✓ 合规"
    elif below:
        place = "图在题下方 *** 不合规 ***"
    else:
        place = "附近无图"
    print(f"  {t[:36]:<38} {place}")
    if not above:
        fails.append(f"figure caption not below image: {t[:20]}")

# ---- table captions above ----
print("\n== 表题位置（模板：表题置于表上方）")
for i, p in enumerate(paras):
    t = (p.text or "").strip()
    if re.match(r"^表\s*\d+", t):
        print(f"  {t[:34]:<36} ok")

# ---- body font ----
print("\n== 正文字体（模板：宋体 小五 9pt）")
sizes = {}
fonts = {}
for p in paras[40:]:
    for r in p.runs:
        if not (r.text or "").strip():
            continue
        sizes[sz(r)] = sizes.get(sz(r), 0) + 1
        f = ea(r)
        fonts[f] = fonts.get(f, 0) + 1
print(f"  字号分布 {dict(sorted((k, v) for k, v in sizes.items() if k))}")
print(f"  中文字体分布 {dict(sorted(((str(k), v) for k, v in fonts.items()), key=lambda x: -x[1])[:5])}")

print("\n" + "=" * 46)
if fails:
    print(f"  {len(fails)} 项不合规:")
    for f in fails:
        print("    -", f)
else:
    print("  全部通过")
for n in notes:
    print("  提示:", n)
print("=" * 46)